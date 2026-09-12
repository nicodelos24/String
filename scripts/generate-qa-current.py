"""Genera las plantillas actuales de String; conserva los Excel históricos Traste.

Requiere openpyxl, solo para documentación. Regenerar sobrescribe las plantillas
String: guardar las ejecuciones manuales personales con otro nombre.
Fuente de casos y bugs: docs/qa/REVISION-ACTUAL.json.
"""
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / 'docs' / 'qa'
DATA = json.loads((QA / 'REVISION-ACTUAL.json').read_text(encoding='utf-8'))

def sheet(book, name, headers, rows, widths):
    ws = book.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='293B33')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F1E8DC')
        ws.row_dimensions[row[0].row].height = 85
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    return ws

def workbook():
    book = Workbook()
    book.remove(book.active)
    sheet(book, 'Leer primero', ['Tema', 'Detalle'], [
        ['Proyecto', 'String, desarrollado con asistencia de IA. Español neutro y ejemplos para aprender QA.'],
        ['Fecha y base', DATA['date'] + ' · ' + DATA['base'] + ' + cambios locales'],
        ['Casos manuales', DATA['manualNote']],
        ['Cómo usar', 'Completar resultado observado, estado, entorno y evidencia después de ejecutar. No confundir cobertura automática con aprobación manual.'],
        ['Regeneración', 'python scripts/generate-qa-current.py. Sobrescribe estas plantillas; guarda tus ejecuciones con otro nombre.'],
        ['Historial', 'Los Excel Traste y las evidencias anteriores se conservan como entregables históricos.'],
        ['Evidencia actual', 'docs/qa/RESULTADOS.md y docs/qa/evidencia/revision-2026-09-11/'],
    ], [24, 100])
    report_path = QA / 'evidencia' / ('revision-' + DATA['date']) / 'ejecucion.json'
    if report_path.exists():
        report = json.loads(report_path.read_text(encoding='utf-8'))
        unit = report['unit']
        rows = [['npm test', f"{unit['pass']}/{unit['tests']} aprobadas", unit['evidence'], 'Pruebas automatizadas; no aprobación manual']]
        rows += [[run['command'], 'Aprobado' if run['exitCode'] == 0 else 'Fallido', run['evidence'], run.get('browser', 'Ver informe')] for run in report['checks']]
        sheet(book, 'Ejecuciones automáticas', ['Comando', 'Resultado', 'Evidencia', 'Entorno / alcance'], rows, [70, 25, 35, 65])
    return book

cases = workbook()
ws = sheet(cases, 'Casos manuales', ['ID', 'Título', 'Precondiciones', 'Pasos', 'Resultado esperado', 'Cobertura automática', 'Bug', 'Estado manual', 'Resultado observado', 'Entorno', 'Evidencia'], [
    [c['id'], c['title'], c['preconditions'], c['steps'], c['expected'], c['automation'], c['bug'], c['status'], c['observed'], '', c['evidence']] for c in DATA['cases']
], [13, 35, 45, 65, 65, 45, 15, 19, 55, 35, 45])
validation = DataValidation(type='list', formula1='"Pendiente,Aprobado,Fallido,Bloqueado"')
ws.add_data_validation(validation)
validation.add(f'H2:H{ws.max_row}')
bugs = workbook()
sheet(bugs, 'Bugs revisión actual', ['ID', 'Título', 'Origen', 'Severidad', 'Prioridad', 'Pasos', 'Esperado', 'Observado', 'Corrección', 'Caso', 'Regresión', 'Estado', 'Evidencia'], [
    [b[k] for k in ['id', 'title', 'origin', 'severity', 'priority', 'steps', 'expected', 'observed', 'fix', 'test', 'automation', 'status', 'evidence']] for b in DATA['bugs']
], [13, 40, 45, 15, 15, 60, 60, 65, 65, 15, 45, 38, 45])
sheet(bugs, 'Historial', ['Registro', 'Ubicación'], [['BUG-001 a BUG-008', 'docs/qa/BUGS.md: se conserva el estado histórico y la aceptación manual pendiente de sonido.']], [35, 100])
out = QA / 'excel'
out.mkdir(exist_ok=True)
for book, filename in [(cases, 'Casos_de_prueba_String.xlsx'), (bugs, 'Reporte_de_bugs_String.xlsx')]:
    path = out / filename
    book.save(path)
    check = load_workbook(path)
    print(f'{filename}: {path.stat().st_size} bytes; hojas verificadas: {check.sheetnames}')
    check.close()

manual = '# Casos de prueba manuales actuales\n\n'
manual += DATA['manualNote'] + '\n\nLos ID MAN-01 a MAN-18 se conservan y actualizan al comportamiento actual. Los casos nuevos son MAN-19 a MAN-38.\n\n'
manual += '| ID | Caso | Pasos | Esperado | Estado manual | Automatización / bug |\n| --- | --- | --- | --- | --- | --- |\n'
for c in DATA['cases']:
    values = [c['id'], c['title'], c['steps'], c['expected'], c['status'], c['automation'] + ('; ' + c['bug'] if c['bug'] else '')]
    manual += '| ' + ' | '.join(v.replace('|', '/') for v in values) + ' |\n'
manual += '\n## Registrar una ejecución\n\nUsa el Excel [Casos de prueba String](excel/Casos_de_prueba_String.xlsx) como plantilla. Completa navegador, versión, dispositivo, pasos realmente ejecutados, resultado observado y evidencia. Guarda una copia con fecha. Si falla, vincula el ID del bug. No marques Aprobado sin ejecutar.\n'
(QA / 'CASOS-MANUALES.md').write_text(manual, encoding='utf-8')
bug_path = QA / 'BUGS.md'
history = bug_path.read_text(encoding='utf-8').split('\n## Revisión 2026-09-11')[0].rstrip()
addition = '\n\n## Revisión 2026-09-11\n\nDefectos encontrados en revisión técnica, no reportes manuales inventados. Los estados siguientes se refieren a la verificación automatizada.\n'
for b in DATA['bugs']:
    addition += f"\n### {b['id']} — {b['title']}\n\n"
    for label, key in [('Origen', 'origin'), ('Severidad', 'severity'), ('Prioridad', 'priority'), ('Pasos', 'steps'), ('Esperado', 'expected'), ('Observado en el código revisado', 'observed'), ('Corrección', 'fix'), ('Caso manual', 'test'), ('Regresión', 'automation'), ('Estado', 'status'), ('Evidencia', 'evidence')]:
        addition += f"- **{label}:** {b[key]}\n"
bug_path.write_text(history + addition, encoding='utf-8')
