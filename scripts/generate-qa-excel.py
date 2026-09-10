"""Genera ejemplos de QA en Excel a partir de la documentación del proyecto.

Requiere openpyxl. Ejecutar desde la raíz: python scripts/generate-qa-excel.py
Las planillas son entregables editables. Regenerar reemplaza los archivos: guardar
las ejecuciones personales con otro nombre antes de volver a usar este script.
"""
from pathlib import Path
import re
import json
import sys
from datetime import datetime, timedelta, timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / 'docs' / 'qa'
OUT = QA / 'excel'
OUT.mkdir(exist_ok=True)
NAVY, TEAL, LIGHT, GRAY = '17324D', '167D8D', 'EAF3F5', '536475'
TODAY = datetime.now(timezone(timedelta(hours=-3))).date().isoformat()
SUFFIX = '_Ritmos' if '--rhythms' in sys.argv else ''


def workbook(title):
    book = Workbook()
    book.remove(book.active)
    book.properties.title = title
    book.properties.creator = 'Traste — documentación de aprendizaje'
    book.properties.description = 'Ejemplo editable. Distingue reportes, pruebas pendientes y verificaciones automatizadas.'
    return book


def heading(ws, title, subtitle, cols):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.cell(1, 1, title).font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    ws.cell(1, 1).fill = PatternFill('solid', fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    ws.cell(2, 1, subtitle).font = Font(name='Calibri', size=11, color=GRAY)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 44
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'Traste | QA | Página &P de &N'


def grid(book, name, title, subtitle, headers, rows, widths):
    ws = book.create_sheet(name)
    heading(ws, title, subtitle, len(headers))
    for column, label in enumerate(headers, 1):
        cell = ws.cell(4, column, label)
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[4].height = 32
    for r, values in enumerate(rows, 5):
        max_lines = 1
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.font = Font(name='Calibri', size=11, color=NAVY)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=Side(style='hair', color='DCE5EC'))
            if r % 2: cell.fill = PatternFill('solid', fgColor='F2F7FA')
            if value:
                max_lines = max(max_lines, sum(max(1, (len(line) + int(widths[c-1]) - 1) // int(widths[c-1])) for line in str(value).split('\n')))
        ws.row_dimensions[r].height = min(230, max(40, max_lines * 16))
    for c, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    end = max(5, ws.max_row)
    table = Table(displayName='Tabla' + re.sub(r'[^A-Za-z0-9]', '', name), ref=f'A4:{get_column_letter(len(headers))}{end}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = 'C5' if len(headers) > 3 else 'B5'
    ws.sheet_view.zoomScale = 85 if len(headers) > 3 else 100
    ws.print_title_rows = '1:4'
    ws.print_options.horizontalCentered = True
    return ws


def choices(ws, column, values):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
    dv.errorTitle = 'Elegí un valor de la lista'
    dv.error = 'Usá una de las opciones disponibles para mantener el mismo criterio.'
    dv.showErrorMessage = True
    dv.errorStyle = 'stop'
    ws.add_data_validation(dv)
    dv.add(f'{column}5:{column}1000')


def color_status(ws, column):
    for state, color in [('Aprobado','DDF0E3'),('Fallido','FADADD'),('Pendiente','FFF1CC'),('Bloqueado','E4E5EB')]:
        ws.conditional_formatting.add(f'{column}5:{column}1000', CellIsRule(operator='equal', formula=[f'"{state}"'], fill=PatternFill('solid',fgColor=color)))


def guide(book, title, specific):
    rows = [
        ('Para qué lo uso', specific),
        ('Cómo está escrito', 'Uso frases simples: qué hago, qué espero y qué pasó. Puedo adaptar el texto a mi forma de expresarme.'),
        ('Qué no debo completar por adelantado', 'No marco Aprobado ni invento evidencia. Primero ejecuto el caso y después escribo el resultado que vi.'),
        ('Origen de los datos', 'Los bugs anteriores vienen de la revisión del código. El problema del reproductor lo reporté durante el uso. Las pruebas hechas por la herramienta están identificadas como automatizadas.'),
        ('Caso de prueba', 'Describe una comprobación que quiero hacer, con datos, pasos y resultado esperado.'),
        ('Ejecución', 'Es una vez que realizo esa comprobación. Si la repito después de una corrección, agrego otra fila para conservar el historial.'),
        ('Severidad', 'Describe cuánto afecta el problema al uso. Por ejemplo, no poder reproducir una progresión tiene impacto alto.'),
        ('Prioridad', 'Indica qué conviene resolver primero. Puede cambiar según el objetivo de la entrega.'),
        ('Entorno', 'Anoto navegador y versión, sistema operativo y versión del proyecto. Si no tengo un dato, escribo Por confirmar.'),
        ('Evidencia', 'Agrego una captura, grabación o registro que realmente exista. Una captura sola no demuestra cómo suena el audio.'),
        ('Planilla editable', 'Hay filtros, encabezados fijos y listas desplegables. Las celdas no están bloqueadas. Para ampliar una tabla, escribo en la fila siguiente.'),
        ('Guardar mi trabajo', 'Guardo una copia con fecha para mis ejecuciones. El script de generación reemplaza estos ejemplos si lo vuelvo a ejecutar.'),
        ('Siguiente paso', 'Recargo la app con Ctrl+F5, pruebo reproducir, repetir y detener. Después completo los casos pendientes, incluido el volumen.'),
    ]
    grid(book,'Leer primero',title,'Ejemplo para aprender QA manual. Lenguaje sencillo, resultados verificables y campos editables.', ['Tema','Explicación'],rows,[28,105])


bug_meta = [
    ('BUG-001','Armadura incorrecta en pentatónica menor','Media','Media','Corregido en código','Revisión anterior', 'Regresión automatizada aprobada. Manual pendiente.'),
    ('BUG-002','Tríada resalta séptimas y novenas','Media','Media','Corregido en código','Revisión anterior', 'Regresión automatizada aprobada. Manual pendiente.'),
    ('BUG-003','La nota no aumenta al pasar el cursor','Baja','Baja','Pendiente de verificación','Revisión anterior', 'Corrección estática. Falta comprobar el efecto visual.'),
    ('BUG-004','El mensaje promete escuchar una nota','Baja','Baja','Pendiente de verificación','Revisión anterior', 'Texto corregido. Falta comprobarlo en el uso.'),
    ('BUG-005','Solo suena un acorde y aparece un error','Alta','Alta','Pendiente de verificación','Reporte del usuario + Chrome real', 'Corregido; comprobación automatizada en navegador. Revalidación manual pendiente.'),
    ('BUG-006','El acorde se escucha muy suave','Media','Media','Pendiente de verificación','Reporte del usuario', 'Se ajustó el sonido. Falta escuchar si el volumen resulta adecuado.'),
    ('BUG-007','No queda claro qué acorde reproduce el botón','Baja','Media','Pendiente de verificación','Reporte del usuario', 'Se aclaró la ayuda. Es una confusión de uso, sin fallo de notas confirmado por separado.'),
    ('BUG-008','Algunas notas suenan más fuertes que otras','Media','Media','Pendiente de verificación','Reporte manual del 2026-09-10', 'Balance ajustado. Falta revalidar con escucha; no se afirma que todas las notas suenen idénticas.'),
]
bug_details = [
    ('1. Elijo A como raíz.\n2. Elijo pentatónica menor.\n3. Reviso la armadura.', 'Espero ver Do mayor como referencia, sin alteraciones.', 'En la revisión anterior se encontró que se mostraba La mayor.', 'Versión anterior a la corrección; navegador no registrado.'),
    ('1. Selecciono Fm7 en las tarjetas.\n2. Activo el resaltado de tríada.\n3. Reviso qué intervalos quedan resaltados.', 'Espero ver raíz, tercera y quinta.', 'En la revisión anterior se encontró que también se resaltaba la séptima.', 'Versión anterior a la corrección; navegador no registrado.'),
    ('1. Abro la página.\n2. Paso el cursor sobre una nota del mástil.', 'Espero que la nota aumente un poco de tamaño.', 'La revisión del CSS encontró transformm en lugar de transform. No hay ejecución visual registrada.', 'Detectado por revisión de código.'),
    ('1. Leo el mensaje debajo del mástil.\n2. Hago clic en una nota.', 'Espero que el mensaje describa lo que hace el clic.', 'El mensaje anterior decía escuchar, pero el clic solo mostraba información.', 'Detectado en revisión; ejecución manual pendiente.'),
    ('1. Abro index.html con doble clic.\n2. Dejo varias tarjetas.\n3. Activo Repetir.\n4. Presiono Reproducir.\n5. Espero el siguiente acorde.', 'Espero escuchar los acordes en orden y que vuelvan a empezar cuando termina la vuelta.', 'Escucho un acorde suave una sola vez. Aparece: No se pudo iniciar el audio. Volvé a intentar.', 'Usuario: archivo local, Chrome o Edge; versión por confirmar. Reproducción técnica: Chrome 152, headless, file://.'),
    ('1. Dejo el volumen inicial.\n2. Presiono Reproducir.\n3. Escucho el nivel y cuánto dura el acorde.', 'Espero escuchar el acorde durante el compás y poder regular el volumen.', 'El acorde se escucha muy suave. Todavía tengo que volver a probarlo con la corrección.', 'Reporte del usuario. Altavoces/auriculares y volumen del sistema por confirmar.'),
    ('1. Cambio la nota o calidad en el selector.\n2. No uso añadir acorde.\n3. Presiono Reproducir.', 'Al usarlo pensé que escucharía la selección. La función definida reproduce las tarjetas guardadas desde la primera.', 'Reporté que no cambiaba el acorde. El fallo de reproducción también impedía avanzar. Falta comprobar si la nueva ayuda se entiende.', 'Reporte del usuario; secuencia exacta de selección por confirmar.'),
    ('1. Reproduzco los acordes.\n2. Escucho el balance entre las notas.\n3. Para la nueva prueba anoto qué acorde o nota sobresale.', 'Espero un balance cómodo entre notas, sin diferencias molestas.', 'Ahora se reproduce sonido. Suena lindo, pero hay notas que se escuchan muy fuertes y otras no tanto.', 'Reporte exploratorio del usuario; raíces, BPM y salida de audio por confirmar.'),
]

bugs = workbook('Traste — Reporte de bugs')
guide(bugs,'Reporte de bugs · Traste','Registro problemas de forma clara para poder reproducirlos, corregirlos y volver a probarlos.')
summary = grid(bugs,'Registro','Registro de bugs','No se cerraron como verificados manualmente los bugs que todavía necesitan una nueva prueba del usuario.',
    ['ID','Título','Severidad','Prioridad','Estado','Origen','Verificación'],bug_meta,[14,40,14,14,28,32,64])
choices(summary,'C',['Alta','Media','Baja'])
choices(summary,'D',['Alta','Media','Baja'])
choices(summary,'E',['Abierto','En corrección','Corregido en código','Pendiente de verificación','Cerrado'])
raw_bugs = (QA/'BUGS.md').read_text(encoding='utf-8')
for number, (meta, details) in enumerate(zip(bug_meta,bug_details),1):
    identifier,title,severity,priority,state,origin,verification = meta
    steps,expected,observed,environment = details
    source_section = re.search(r'## '+identifier+r'[^\n]*\n(.*?)(?=\n## |\Z)',raw_bugs,re.S).group(1)
    technical = []
    for label, value in re.findall(r'- \*\*([^*]+):\*\* (.+)',source_section):
        if label in ['Causa','Análisis','Solución','Solución aplicada','Evidencia','Regresión','Caso vinculado']:
            technical.append((label,value))
    rows = [('ID',identifier),('Título',title),('Origen',origin),('Fecha de actualización',TODAY),
        ('Estado',state),('Severidad',severity),('Prioridad',priority),('Entorno',environment),
        ('Precondiciones','Tengo la página abierta. Para audio, dejo el metrónomo detenido.'),('Pasos',steps),
        ('Resultado esperado',expected),('Resultado observado',observed),('Frecuencia','Por confirmar; no se registró un número de intentos del usuario.'),
        ('Verificación',verification),*technical,('Nueva prueba manual','Pendiente. Completo después de probar: fecha, pasos, resultado y evidencia.')]
    grid(bugs,identifier,identifier+' · '+title,'Los campos técnicos vienen de la investigación. El resultado manual posterior a la corrección sigue pendiente.', ['Campo','Detalle'],rows,[28,115])
    summary.cell(number+4,1).hyperlink = f"#'{identifier}'!A1"
    summary.cell(number+4,1).font = Font(color=TEAL,underline='single')
grid(bugs,'Plantilla','Mi próximo reporte','Duplico esta hoja para reportar un problema nuevo. Completo lo que realmente observé.', ['Campo','Detalle'],
    [(field,value) for field,value in [('ID','BUG-009'),('Título',''),('Fecha',''),('Entorno',''),('Precondiciones',''),('Pasos','1.\n2.\n3.'),('Resultado esperado',''),('Resultado observado',''),('Severidad',''),('Prioridad',''),('Estado','Abierto'),('Evidencia',''),('Caso relacionado',''),('Resultado de la nueva prueba','Pendiente')]], [28,115])
bugs.save(OUT/f'Reporte_de_bugs_Traste{SUFFIX}.xlsx')


cases_book = workbook('Traste — Casos de prueba')
guide(cases_book,'Casos de prueba · Traste','Organizo las comprobaciones que voy a hacer. Registro cada ejecución en otra hoja, para no mezclar lo esperado con lo observado.')
manual_text = (QA/'CASOS-MANUALES.md').read_text(encoding='utf-8').split('## Plantilla')[0]
source_rows = [line.strip('| ').split(' | ') for line in manual_text.splitlines() if line.startswith('| MAN-')]
titles = ['Reproducir la progresión inicial','Comprobar duración a 120 BPM','Terminar una vuelta y repetir','Detener y volver a iniciar',
          'Validar límites de BPM','Cambiar volumen durante el acorde','Editar tarjetas mientras suena','Recuperarse de un error de audio',
          'Salir de la página y volver','Usar los controles con teclado','Usar la app en móvil y escritorio','Revisar funciones anteriores','Añadir un acorde elegido al reproductor',
          'Comparar estilos','Desactivar percusión','Regular la mezcla','Repetir y detener con batería','Comparar balance de notas']
auto = ['PLY-01/02/12 + navegador','PLY-02','PLY-02/03 + navegador','PLY-04/05/11/14 + navegador','PLY-06/13',
        'PLY-08/15 (sin escucha)','PLY-07/12','PLY-10/13','PLY-09/13 (parcial)','Pendiente','Pendiente','Suite anterior (parcial)','Navegador: progresión distinta (parcial)',
        'RIT-01/02/06 + navegador','RIT-03','RIT-04','RIT-03 + motor base (parcial)','MIX-01 + OfflineAudioContext; escucha pendiente']
rows = []
for i,(identifier,steps,expected) in enumerate(source_rows):
    # Numerar los pasos facilita repetir el caso sin adoptar un tono artificial.
    split_steps = re.split(r'\. (?=[A-ZÁÉÍÓÚ])',steps.rstrip('.'))
    steps = '\n'.join(f'{n}. {step.rstrip(".")}.' for n,step in enumerate(split_steps,1))
    rows.append([identifier,titles[i],'Página de raíz abierta; metrónomo detenido. Recargo antes del caso.',
        'C, Fm7, G7, C; salvo indicación en los pasos.',steps,expected,'Alta' if i in [0,1,2,3,4] else 'Media',auto[i],'Pendiente'])
case_sheet = grid(cases_book,'Casos','Casos de prueba manuales','Cada caso tiene su esperado. La columna Estado manual sigue pendiente; el reporte exploratorio está en Ejecuciones.',
    ['ID','Objetivo','Precondiciones','Datos','Pasos','Resultado esperado','Prioridad','Automatización relacionada','Estado manual'],rows,[14,34,38,33,67,76,14,38,20])
choices(case_sheet,'G',['Alta','Media','Baja'])
choices(case_sheet,'I',['Pendiente','Aprobado','Fallido','Bloqueado'])
color_status(case_sheet,'I')

execution_rows = [[
    'EXP-001','Exploratoria; no fue una ejecución formal de MAN-01','2026-09-09','Manual: reporte del usuario',
    'Doble clic en index.html. Chrome o Edge; versión exacta por confirmar.',
    'Presiono Reproducir y suena un acorde suave una vez; no avanza ni repite y aparece un error. También reporté que no cambia al elegir otro acorde.',
    'Fallido','BUG-005 / BUG-006 / BUG-007','Reporte escrito en la conversación. Sin captura ni grabación adjunta.'
],[
    'EXP-002-A','Solo inicio de sonido; confirmación parcial','2026-09-10','Manual: reporte del usuario',
    'Mismo proyecto. Navegador y configuración exactos por confirmar.',
    'Ahora el sonido se reproduce. No confirmé por separado todos los casos de repetición, detención ni los nuevos estilos.',
    'Aprobado','BUG-005 (solo inicio de audio)','Reporte escrito del usuario; sin grabación.'
],[
    'EXP-002-B','Balance entre notas','2026-09-10','Manual: reporte del usuario',
    'Raíces, BPM y dispositivo de salida por confirmar.',
    'Suena lindo pero algunas notas se escuchan muy fuertes y otras no tanto. Falta probar el ajuste nuevo.',
    'Fallido','BUG-008','Reporte escrito del usuario; sin grabación.'
]]
for filename,label in [('reproductor-antes.json','Antes de la corrección'),('reproductor-despues.json','Después, Chrome'),('reproductor-despues-edge.json','Después, Edge')]:
    evidence = QA/'evidencia'/filename
    if not evidence.exists(): continue
    data = json.loads(evidence.read_text(encoding='utf-8'))
    passed = data['diagnostic'] == 'OK' and sum('acorde 1 de 4' in state for state in data['playback']['states']) >= 2 and data['stopped'] == 'Detenido'
    execution_rows.append([f'AUT-{len(execution_rows):03d}','MAN-01/03/04: parte automatizada',data['date'],'Automatizada: navegador real',
        data['browser']+'; '+data['mode'],label+': '+('avanza, repite y se detiene; sin escucha humana.' if passed else 'solo se programa el primer acorde y se registra Illegal invocation.'),
        'Aprobado' if passed else 'Fallido','BUG-005','../evidencia/'+filename])
for filename in ['reproductor-ritmos-chrome.json','reproductor-ritmos-edge.json']:
    evidence = QA/'evidencia'/filename
    if not evidence.exists(): continue
    data = json.loads(evidence.read_text(encoding='utf-8'))
    passed = len(data.get('rhythms',[])) == 3 and all(item['status']=='Detenido' for item in data['rhythms'])
    execution_rows.append([f'RIT-AUT-{len(execution_rows)}','MAN-14/17/18: parte automatizada',data['date'],
        'Automatizada: navegador y audio offline',data['browser'],
        'Tres estilos generan percusión y acordes. Se midieron 12 raíces y 3 patrones. Sin escucha humana.',
        'Aprobado' if passed else 'Fallido','BUG-008: medición técnica, no aceptación auditiva','../evidencia/'+filename])
for identifier,*_ in source_rows:
    execution_rows.append(['',identifier,'','Manual','','','Pendiente','',''])
execution = grid(cases_book,'Ejecuciones','Historial de ejecuciones','Conservo el reporte inicial y separo las comprobaciones automáticas. Completo las filas manuales al probar la corrección.',
    ['Ejecución','Caso / alcance','Fecha','Tipo / origen','Entorno','Resultado observado','Estado','Bug relacionado','Evidencia'],execution_rows,[16,39,27,32,50,80,16,28,57])
choices(execution,'G',['Pendiente','Aprobado','Fallido','Bloqueado'])
color_status(execution,'G')
for row in range(5,execution.max_row+1):
    value = execution.cell(row,9).value
    if value and value.startswith('../evidencia/'):
        execution.cell(row,9).hyperlink = value
        execution.cell(row,9).font = Font(color=TEAL,underline='single')

grid(cases_book,'Cobertura','Qué prueban las automatizaciones','34 pruebas de Node aprobadas. Las comprobaciones de navegador se ejecutan con otro comando y no cuentan dentro de esas 34.',
    ['Grupo','Comando / archivo','Qué comprueba','Límite'],[
        ['Metrónomo','tests/metronome.test.cjs','Tempo, acento, detención e inicio cancelado.','Audio simulado.'],
        ['Progresión y controles','tests/progression.test.cjs','Selección, escritura, armadura, tríadas y conexión de controles.','DOM simulado.'],
        ['Motor del reproductor','tests/player.test.cjs','Notas, tiempos, repetición, volumen, limpieza y contexto de temporizadores.','No verifica la percepción del volumen.'],
        ['Suite Node','node --test tests/*.test.cjs','34 aprobadas, 0 fallidas.','No equivale a 100% de cobertura ni reemplaza las pruebas manuales.'],
        ['Ritmos y mezcla','RIT-01 a RIT-06; MIX-01','Patrones, tempo, percusión independiente y normalización de ganancias.','Aceptación de estilo y sonoridad pendiente.'],
        ['Chrome/Edge real','node scripts/check-player-browser.cjs [--edge]','Archivo local, cambio de acordes, repetición, detención y secuencia distinta.','Headless y audio silenciado: no es una escucha manual.'],
    ],[27,65,72,68])
cases_book.save(OUT/f'Casos_de_prueba_Traste{SUFFIX}.xlsx')

# Verificación estructural: abrir de nuevo cada archivo, revisar contenido y reglas.
for filename in [f'Reporte_de_bugs_Traste{SUFFIX}.xlsx',f'Casos_de_prueba_Traste{SUFFIX}.xlsx']:
    output = OUT/filename
    reopened = load_workbook(output)
    assert len(reopened.sheetnames) >= 4
    assert all(ws.freeze_panes for ws in reopened)
    assert all(ws.tables for ws in reopened)
    assert not any(cell.data_type == 'f' for ws in reopened for row in ws for cell in row)
    print(f'{filename}: {len(reopened.sheetnames)} hojas; {output.stat().st_size} bytes; reabierto correctamente.')
